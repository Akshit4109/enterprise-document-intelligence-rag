import io
import re
from typing import Any, Dict, List
import docx
from app.core.exceptions import CorruptedFileException, ParserException
from app.parsers.base import BaseDocumentParser, ParsedDocument, ParsedPage


class DOCXParser(BaseDocumentParser):
    """
    Parser for Microsoft Word (.docx) documents using python-docx.
    Extracts text from paragraphs and embedded tables while capturing core document properties.
    """

    def parse(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            raise CorruptedFileException(
                message=f"Failed to open DOCX document '{filename}'. File may be corrupt or invalid.",
                details={"filename": filename, "original_error": str(e)}
            )

        try:
            paragraphs_text: List[str] = []
            for paragraph in doc.paragraphs:
                p_text = paragraph.text.strip()
                if p_text:
                    paragraphs_text.append(p_text)

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs_text.append(row_text)

            full_text = "\n\n".join(paragraphs_text).strip()
            full_text = self._clean_text(full_text)

            # Metadata extraction
            doc_metadata: Dict[str, Any] = {"format": "DOCX"}
            if hasattr(doc, "core_properties"):
                core_props = doc.core_properties
                doc_metadata.update({
                    "title": core_props.title or "",
                    "author": core_props.author or "",
                    "subject": core_props.subject or "",
                    "category": core_props.category or "",
                    "comments": core_props.comments or "",
                })

            # DOCX does not have explicit fixed pagination in raw XML; represent as 1 section or by headings
            pages: List[ParsedPage] = [
                ParsedPage(
                    page_number=1,
                    text=full_text,
                    metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)}
                )
            ]

            words = full_text.split()
            word_count = len(words)
            char_count = len(full_text)

            return ParsedDocument(
                text=full_text,
                pages=pages,
                total_pages=1,
                word_count=word_count,
                character_count=char_count,
                metadata=doc_metadata,
            )
        except Exception as e:
            raise ParserException(
                message=f"Error parsing DOCX '{filename}': {str(e)}",
                details={"filename": filename, "original_error": str(e)}
            )

    @staticmethod
    def _clean_text(text: str) -> str:
        text = re.sub(r"\r\n|\r", "\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
