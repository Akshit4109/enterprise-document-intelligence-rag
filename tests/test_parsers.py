import io
import fitz  # PyMuPDF
import docx
import pytest

from app.core.exceptions import (
    CorruptedFileException,
    UnsupportedFileFormatException,
)
from app.parsers.factory import DocumentParserFactory
from app.parsers.pdf import PDFParser
from app.parsers.docx import DOCXParser
from app.parsers.txt import TXTParser


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    """Generates a real multi-page PDF in-memory using PyMuPDF."""
    doc = fitz.open()
    
    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Enterprise Document Intelligence Platform.\nPage 1 Header and Introduction.")
    
    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Security and Architecture Specifications.\nPage 2 Implementation Details.")
    
    doc.set_metadata({"title": "System Architecture", "author": "Enterprise Architect"})
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Generates a real DOCX file in-memory using python-docx."""
    doc = docx.Document()
    doc.core_properties.title = "Employee Handbook"
    doc.core_properties.author = "HR Department"
    
    doc.add_heading("Employee Conduct Policy", level=1)
    doc.add_paragraph("All employees must adhere to the high standard of data security and ethics.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Access Level"
    table.cell(1, 0).text = "Engineer"
    table.cell(1, 1).text = "Read/Write"
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


@pytest.fixture
def sample_txt_bytes() -> bytes:
    """Generates sample plain text bytes."""
    return "Enterprise Platform README\n\nOverview:\nThis system handles scalable RAG document ingestion.".encode("utf-8")


class TestPDFParser:
    def test_parse_valid_pdf_multi_page(self, sample_pdf_bytes: bytes):
        parser = PDFParser()
        result = parser.parse(sample_pdf_bytes, "architecture.pdf")

        assert result.total_pages == 2
        assert len(result.pages) == 2
        assert result.pages[0].page_number == 1
        assert "Enterprise Document Intelligence" in result.pages[0].text
        assert result.pages[1].page_number == 2
        assert "Security and Architecture" in result.pages[1].text
        assert result.metadata["title"] == "System Architecture"
        assert result.word_count > 0
        assert result.character_count > 0

    def test_parse_corrupted_pdf(self):
        parser = PDFParser()
        with pytest.raises(CorruptedFileException):
            parser.parse(b"NOT_A_REAL_PDF_STREAM", "corrupted.pdf")


class TestDOCXParser:
    def test_parse_valid_docx(self, sample_docx_bytes: bytes):
        parser = DOCXParser()
        result = parser.parse(sample_docx_bytes, "handbook.docx")

        assert result.total_pages == 1
        assert "Employee Conduct Policy" in result.text
        assert "Engineer | Read/Write" in result.text
        assert result.metadata["title"] == "Employee Handbook"
        assert result.metadata["author"] == "HR Department"
        assert result.word_count > 0

    def test_parse_corrupted_docx(self):
        parser = DOCXParser()
        with pytest.raises(CorruptedFileException):
            parser.parse(b"CORRUPT_DOCX_BYTES", "broken.docx")


class TestTXTParser:
    def test_parse_valid_txt(self, sample_txt_bytes: bytes):
        parser = TXTParser()
        result = parser.parse(sample_txt_bytes, "readme.txt")

        assert result.total_pages == 1
        assert "Enterprise Platform README" in result.text
        assert result.metadata["format"] == "TXT"
        assert result.word_count > 0

    def test_parse_latin1_encoded_txt(self):
        parser = TXTParser()
        content = "Café résumé naïve".encode("latin-1")
        result = parser.parse(content, "french.txt")

        assert "Café résumé naïve" in result.text


class TestParserFactory:
    def test_factory_resolves_correct_parsers(self):
        pdf_parser = DocumentParserFactory.get_parser("document.pdf")
        docx_parser = DocumentParserFactory.get_parser("doc.docx")
        txt_parser = DocumentParserFactory.get_parser("notes.txt")

        assert isinstance(pdf_parser, PDFParser)
        assert isinstance(docx_parser, DOCXParser)
        assert isinstance(txt_parser, TXTParser)

    def test_factory_unsupported_format(self):
        with pytest.raises(UnsupportedFileFormatException):
            DocumentParserFactory.get_parser("script.exe")
