import io
import fitz
import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from sqlalchemy import select

from app.main import app
from app.core.database import get_db
from app.models.document import Document
from app.models.document_version import DocumentVersion


@pytest.fixture
def client(db_session: Session) -> TestClient:
    """FastAPI TestClient with overridden get_db dependency for transactional testing."""
    def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "FastAPI Upload Test PDF Content.")
    doc.set_metadata({"title": "API Test Document"})
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("API Upload Test DOCX", level=1)
    doc.add_paragraph("Testing python-docx upload endpoint.")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


class TestUploadAPI:
    def test_upload_pdf_success(self, client: TestClient, db_session: Session, sample_pdf_bytes: bytes):
        files = {"file": ("test_doc.pdf", sample_pdf_bytes, "application/pdf")}
        data = {
            "name": "Integration Test Document",
            "description": "Uploaded via pytest",
            "owner_id": "usr_test_101",
        }

        response = client.post("/api/v1/documents/upload", files=files, data=data)
        assert response.status_code == 201
        res_json = response.json()

        assert res_json["message"] == "Document uploaded, parsed, chunked, and embedded successfully"
        assert res_json["document"]["name"] == "Integration Test Document"
        assert res_json["document"]["document_type"] == "pdf"
        assert res_json["version"]["version_number"] == 1
        assert res_json["parsing_summary"]["total_pages"] == 1
        assert res_json["parsing_summary"]["chunk_count"] > 0
        assert "FastAPI Upload Test PDF" in res_json["parsing_summary"]["preview_text"]

        # Verify DB records
        doc_id = res_json["document"]["id"]
        version_id = res_json["version"]["id"]
        doc = db_session.get(Document, doc_id)
        assert doc is not None
        assert len(doc.versions) == 1
        assert doc.versions[0].version_number == 1
        assert doc.versions[0].file_name == "test_doc.pdf"
        assert len(doc.versions[0].chunks) > 0

        # Verify GET /documents/{id}/versions/{version_id}/chunks
        chunks_resp = client.get(f"/api/v1/documents/{doc_id}/versions/{version_id}/chunks")
        assert chunks_resp.status_code == 200
        chunks_json = chunks_resp.json()
        assert len(chunks_json) == len(doc.versions[0].chunks)
        assert chunks_json[0]["chunk_index"] == 0

    def test_upload_docx_success(self, client: TestClient, db_session: Session, sample_docx_bytes: bytes):
        files = {"file": ("sample.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {"owner_id": "usr_test_102"}

        response = client.post("/api/v1/documents/upload", files=files, data=data)
        assert response.status_code == 201
        res_json = response.json()

        assert res_json["document"]["document_type"] == "docx"
        assert "API Upload Test DOCX" in res_json["parsing_summary"]["preview_text"]

    def test_upload_txt_success(self, client: TestClient, db_session: Session):
        txt_content = b"Plain text file content for ingestion testing."
        files = {"file": ("notes.txt", txt_content, "text/plain")}

        response = client.post("/api/v1/documents/upload", files=files)
        assert response.status_code == 201
        res_json = response.json()

        assert res_json["document"]["document_type"] == "txt"
        assert "Plain text file content" in res_json["parsing_summary"]["preview_text"]

    def test_upload_multiple_versions(self, client: TestClient, db_session: Session, sample_pdf_bytes: bytes):
        # Upload version 1
        files1 = {"file": ("policy_v1.pdf", sample_pdf_bytes, "application/pdf")}
        resp1 = client.post("/api/v1/documents/upload", files=files1, data={"name": "Company Policy"})
        assert resp1.status_code == 201
        doc_id = resp1.json()["document"]["id"]

        # Upload version 2 to existing document_id
        files2 = {"file": ("policy_v2.pdf", sample_pdf_bytes, "application/pdf")}
        resp2 = client.post(
            "/api/v1/documents/upload",
            files=files2,
            data={"document_id": doc_id}
        )
        assert resp2.status_code == 201
        res2_json = resp2.json()

        assert res2_json["document"]["id"] == doc_id
        assert res2_json["version"]["version_number"] == 2

        # Verify through GET /documents/{id}
        get_resp = client.get(f"/api/v1/documents/{doc_id}")
        assert get_resp.status_code == 200
        get_json = get_resp.json()
        assert len(get_json["versions"]) == 2
        assert [v["version_number"] for v in get_json["versions"]] == [2, 1]

    def test_upload_empty_file(self, client: TestClient):
        files = {"file": ("empty.txt", b"", "text/plain")}
        response = client.post("/api/v1/documents/upload", files=files)
        assert response.status_code == 400
        assert "empty" in response.json()["message"].lower()

    def test_upload_unsupported_format(self, client: TestClient):
        files = {"file": ("malicious.exe", b"MZ\x90\x00BinaryData", "application/octet-stream")}
        response = client.post("/api/v1/documents/upload", files=files)
        assert response.status_code == 415
        assert "not supported" in response.json()["message"].lower()

    def test_upload_spoofed_pdf_magic_bytes_failure(self, client: TestClient):
        # File is named .pdf but contains text/binary without %PDF- header
        files = {"file": ("fake.pdf", b"This is plain text pretending to be a PDF.", "application/pdf")}
        response = client.post("/api/v1/documents/upload", files=files)
        assert response.status_code == 415
        assert "magic header" in response.json()["message"].lower()
